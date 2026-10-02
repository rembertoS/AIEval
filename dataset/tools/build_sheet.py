"""Write an xlsx summary of the dev dataset, straight from ground_truth.json.

Usage:
    .venv/bin/python tools/build_sheet.py                 # writes dataset_sheet.xlsx
    .venv/bin/python tools/build_sheet.py --out /tmp/x.xlsx
    .venv/bin/python tools/build_sheet.py --csv-dir sheets/   # csv instead of xlsx

Sheets: Cases, Vulnerabilities, Decoys, Composition.
"""

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PRIMARY_CWES = ("CWE-89", "CWE-78", "CWE-22", "CWE-327")
SUBTLETY_ORDER = {"obvious": 0, "moderate": 1, "subtle": 2}

CASE_COLUMNS = [
    "case_id", "category", "vulns", "decoys", "cwes", "subtleties",
    "max_fix_risk", "app_lines", "context",
]
VULN_COLUMNS = [
    "case_id", "vuln_id", "cwe", "subtlety", "expected_fix_risk", "function",
    "line_start", "line_end", "related_lines", "security_test",
    "source", "sink", "description", "regression_note",
]
DECOY_COLUMNS = [
    "case_id", "decoy_id", "looks_like", "function", "line_start", "line_end",
    "related_lines", "security_test", "why_safe",
]


def _related(entry):
    parts = [
        f"{loc.get('function')}:{loc.get('line_start')}-{loc.get('line_end')}"
        for loc in entry.get("related_lines", [])
    ]
    return ", ".join(parts)


def _app_lines(case_id):
    path = ROOT / case_id / "app.py"
    return len(path.read_text().splitlines()) if path.exists() else 0


def case_rows(cases):
    rows = []
    for case in cases:
        vulns = case.get("vulnerabilities", [])
        risks = [v.get("expected_fix_risk", "") for v in vulns]
        subtleties = sorted(
            (v.get("subtlety", "") for v in vulns),
            key=lambda s: SUBTLETY_ORDER.get(s, 9),
        )
        rows.append({
            "case_id": case["case_id"],
            "category": case["category"],
            "vulns": len(vulns),
            "decoys": len(case.get("decoys", [])),
            "cwes": ", ".join(v["cwe"] for v in vulns),
            "subtleties": ", ".join(subtleties),
            "max_fix_risk": max(risks, key=lambda r: ["Low", "Medium", "High"].index(r)) if risks else "",
            "app_lines": _app_lines(case["case_id"]),
            "context": case.get("context", ""),
        })
    return rows


def vuln_rows(cases):
    rows = []
    for case in cases:
        for vuln in case.get("vulnerabilities", []):
            row = {key: vuln.get(key, "") for key in VULN_COLUMNS}
            row["case_id"] = case["case_id"]
            row["related_lines"] = _related(vuln)
            rows.append(row)
    return rows


def decoy_rows(cases):
    rows = []
    for case in cases:
        for decoy in case.get("decoys", []):
            row = {key: decoy.get(key, "") for key in DECOY_COLUMNS}
            row["case_id"] = case["case_id"]
            row["related_lines"] = _related(decoy)
            rows.append(row)
    return rows


def composition_rows(cases):
    categories = Counter(c["category"] for c in cases)
    vulns = [v for c in cases for v in c.get("vulnerabilities", [])]
    singles = Counter(
        c["vulnerabilities"][0]["cwe"]
        for c in cases
        if c["category"] == "single" and c.get("vulnerabilities")
    )
    cwe798_multi = sum(
        1
        for c in cases
        if c["category"] == "multi"
        and any(v["cwe"] == "CWE-798" for v in c["vulnerabilities"])
    )
    decoyed = sum(
        1 for c in cases if c["category"] != "clean" and c.get("decoys")
    )
    subtleties = Counter(v.get("subtlety") for v in vulns)
    rows = [
        {"rule": "cases", "target": "30", "actual": str(len(cases))},
        {"rule": "single / multi / clean", "target": "16 / 8 / 6",
         "actual": f"{categories['single']} / {categories['multi']} / {categories['clean']}"},
        {"rule": "real vulnerabilities", "target": "34", "actual": str(len(vulns))},
        {"rule": "decoys", "target": "-",
         "actual": str(sum(len(c.get('decoys', [])) for c in cases))},
    ]
    for cwe in PRIMARY_CWES:
        rows.append({"rule": f"singles with {cwe}", "target": "4", "actual": str(singles[cwe])})
    rows += [
        {"rule": "multi cases with CWE-798", "target": "~5", "actual": str(cwe798_multi)},
        {"rule": "single/multi cases with a decoy", "target": ">= 8", "actual": str(decoyed)},
        {"rule": "subtlety obvious / moderate / subtle", "target": "-",
         "actual": f"{subtleties['obvious']} / {subtleties['moderate']} / {subtleties['subtle']}"},
    ]
    for cwe in PRIMARY_CWES + ("CWE-798",):
        per_cwe = [v for v in vulns if v["cwe"] == cwe]
        counts = Counter(v.get("subtlety") for v in per_cwe)
        rows.append({
            "rule": f"{cwe} O/M/S (total)",
            "target": "-",
            "actual": f"{counts['obvious']}/{counts['moderate']}/{counts['subtle']} ({len(per_cwe)})",
        })
    return rows


SHEETS = [
    ("Cases", CASE_COLUMNS, case_rows),
    ("Vulnerabilities", VULN_COLUMNS, vuln_rows),
    ("Decoys", DECOY_COLUMNS, decoy_rows),
    ("Composition", ["rule", "target", "actual"], composition_rows),
]


def write_xlsx(cases, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)
    for title, columns, builder in SHEETS:
        sheet = book.create_sheet(title)
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in builder(cases):
            sheet.append([row.get(column, "") for column in columns])
        sheet.freeze_panes = "A2"
        for index, column in enumerate(columns, start=1):
            widest = max(
                [len(column)] + [len(str(r.get(column, ""))) for r in builder(cases)]
            )
            sheet.column_dimensions[get_column_letter(index)].width = min(60, widest + 2)
            if column in ("context", "description", "why_safe", "regression_note", "source", "sink"):
                for cell in sheet[get_column_letter(index)]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
    book.save(out_path)
    return out_path


def write_csvs(cases, out_dir):
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for title, columns, builder in SHEETS:
        path = out_dir / f"{title.lower()}.csv"
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=columns)
            writer.writeheader()
            writer.writerows(builder(cases))
        written.append(path)
    return written


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", default=str(ROOT / "dataset_sheet.xlsx"))
    parser.add_argument("--csv-dir", help="write csv files instead of an xlsx workbook")
    args = parser.parse_args(argv)

    cases = json.loads((ROOT / "ground_truth.json").read_text())["cases"]

    if args.csv_dir:
        for path in write_csvs(cases, Path(args.csv_dir)):
            print(f"wrote {path}")
    else:
        print(f"wrote {write_xlsx(cases, args.out)}")

    for row in composition_rows(cases):
        print(f"  {row['rule']:<38} target {row['target']:<12} actual {row['actual']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
